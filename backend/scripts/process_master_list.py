"""Process ULTRA_GRANULAR_MASTER_LIST.xlsx:
  1) Dedupe by normalized item name (keep first occurrence with most fields filled).
  2) Run each product through the canonical classifier.
  3) Output:
       /tmp/master_dedup_classified.csv   — full deduped + classified table
       /tmp/master_summary.json           — distribution by niche/category/subcategory
       /tmp/master_unclassified.csv       — rows where category=None (admin curation list)
"""
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from services.taxonomy_canonical import classify_product, _norm_name  # noqa: E402


SRC = "/tmp/master.xlsx"
OUT_FULL = "/tmp/master_dedup_classified.csv"
OUT_SUMMARY = "/tmp/master_summary.json"
OUT_UNCLASSIFIED = "/tmp/master_unclassified.csv"


def _row_completeness(row: tuple) -> int:
    """Higher score = more non-empty cells. Used to pick the best dup record."""
    return sum(1 for v in row if v not in (None, "", "General"))


def main():
    wb = load_workbook(SRC, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    data = rows[1:]
    print(f"Loaded {len(data)} rows ({len(header)} cols)")

    # 1) Dedupe by normalized name. Among duplicates, keep the most-complete row.
    by_key: dict[str, tuple] = {}
    dup_count = 0
    for r in data:
        if not r or not r[1]:
            continue
        key = _norm_name(str(r[1]))
        if not key:
            continue
        if key in by_key:
            dup_count += 1
            if _row_completeness(r) > _row_completeness(by_key[key]):
                by_key[key] = r
        else:
            by_key[key] = r
    print(f"De-duplicated: {len(by_key):,} unique | {dup_count:,} dupes removed")

    # 2) Classify each unique product through canonical classifier.
    out_rows: list[dict] = []
    niche_counter = Counter()
    cat_counter = Counter()
    sub_counter = Counter()
    by_cat_subs = defaultdict(Counter)
    unclassified: list[dict] = []

    for r in by_key.values():
        brand = (r[0] or "").strip()
        name = (r[1] or "").strip()
        main_cat = (r[2] or "").strip()
        ptype = (r[3] or "").strip()
        concern = (r[4] or "").strip()

        result = classify_product(
            name=name,
            description="",
            brand=brand,
            current_niche=main_cat.lower() if main_cat else None,
        )
        row_out = {
            "brand": brand,
            "name": name,
            "src_main_category": main_cat,
            "src_product_type": ptype,
            "src_concern": concern,
            "niche": result["niche"],
            "category": result["category"] or "",
            "subcategory": result["subcategory"] or "",
            "concerns": ",".join(result.get("concerns") or []),
            "needs_review": str(result.get("unclassified", False)).lower(),
            "mrp": r[7] if len(r) > 7 else "",
            "dealer_price": r[8] if len(r) > 8 else "",
            "list_price": r[10] if len(r) > 10 else "",
            "market_segment": (r[6] or "").strip() if len(r) > 6 else "",
        }
        out_rows.append(row_out)
        niche_counter[result["niche"]] += 1
        cat_counter[result["category"] or "__null__"] += 1
        if result["subcategory"]:
            sub_counter[result["subcategory"]] += 1
            by_cat_subs[result["category"]][result["subcategory"]] += 1
        else:
            by_cat_subs[result["category"] or "__null__"]["__no_sub__"] += 1
        if not result["category"]:
            unclassified.append(row_out)

    # 3) Write full deduped CSV
    fieldnames = list(out_rows[0].keys())
    with open(OUT_FULL, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(out_rows)
    print(f"Wrote {OUT_FULL} ({len(out_rows):,} rows)")

    # 4) Write unclassified subset
    if unclassified:
        with open(OUT_UNCLASSIFIED, "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fieldnames)
            w.writeheader()
            w.writerows(unclassified)
        print(f"Wrote {OUT_UNCLASSIFIED} ({len(unclassified):,} rows)")

    # 5) Summary JSON
    summary = {
        "input_rows": len(data),
        "unique_products": len(out_rows),
        "duplicates_removed": dup_count,
        "by_niche": dict(niche_counter),
        "by_category_top30": dict(cat_counter.most_common(30)),
        "by_subcategory_top60": dict(sub_counter.most_common(60)),
        "by_category_with_subs": {
            cat: dict(subs.most_common(20))
            for cat, subs in sorted(by_cat_subs.items(),
                                     key=lambda kv: -sum(kv[1].values()))
        },
        "unclassified_count": len(unclassified),
    }
    with open(OUT_SUMMARY, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2)
    print(f"Wrote {OUT_SUMMARY}")

    # Console report
    print("\n=== NICHE DISTRIBUTION ===")
    for k, v in niche_counter.most_common():
        print(f"  {k:15} {v:>6,}")
    print("\n=== CATEGORY DISTRIBUTION ===")
    for k, v in cat_counter.most_common():
        print(f"  {k:25} {v:>6,}")
    print("\n=== TOP 30 SUBCATEGORIES ===")
    for k, v in sub_counter.most_common(30):
        print(f"  {k:30} {v:>6,}")
    print(f"\nUnclassified (category=None): {len(unclassified):,}")


if __name__ == "__main__":
    main()
