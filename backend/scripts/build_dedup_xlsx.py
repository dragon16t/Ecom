"""Generate a deduped XLSX ready for /api/admin/bulk-import/upload.

Output: /app/memory/master_dedup_for_bulk_import.xlsx
Columns match what services/bulk_import_service.py expects (Brand, Item Name,
MRP, Dealer Price, Listing Price). Also includes the classifier's predicted
niche/category/subcategory so an admin can review before triggering AI auto-fill.
"""
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook, Workbook  # noqa: E402

from services.taxonomy_canonical import classify_product, _norm_name  # noqa: E402


def _completeness(row: tuple) -> int:
    return sum(1 for v in row if v not in (None, "", "General"))


def main():
    src = Path("/tmp/master.xlsx")
    if not src.exists():
        print("/tmp/master.xlsx missing — run scripts/process_master_list.py first")
        sys.exit(1)
    wb_in = load_workbook(src, read_only=True)
    ws_in = wb_in["Sheet1"]
    rows = list(ws_in.iter_rows(values_only=True))
    data = rows[1:]

    by_key: dict[str, tuple] = {}
    for r in data:
        if not r or not r[1]:
            continue
        key = _norm_name(str(r[1]))
        if not key:
            continue
        if key in by_key:
            if _completeness(r) > _completeness(by_key[key]):
                by_key[key] = r
        else:
            by_key[key] = r

    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append([
        "Brand", "Item Name", "MRP", "Dealer Price",
        "Listing Price", "Niche (auto)", "Category (auto)",
        "Subcategory (auto)", "Source Product Type", "Concern (auto)",
        "Needs Review",
    ])

    classified = 0
    review = 0
    for r in by_key.values():
        brand = (r[0] or "").strip() if len(r) > 0 else ""
        name = (r[1] or "").strip()
        src_type = (r[3] or "").strip() if len(r) > 3 else ""
        mrp = r[7] if len(r) > 7 else ""
        dealer = r[8] if len(r) > 8 else ""
        listing = r[10] if len(r) > 10 else ""

        res = classify_product(
            name=name, description="", brand=brand,
            current_niche=(r[2] or "").lower() if len(r) > 2 and r[2] else None,
        )
        if res["category"]:
            classified += 1
        else:
            review += 1
        ws.append([
            brand, name, mrp, dealer, listing,
            res["niche"] or "", res["category"] or "",
            res["subcategory"] or "", src_type,
            ",".join(res.get("concerns") or []),
            "YES" if not res["category"] else "",
        ])

    out = Path("/app/memory/master_dedup_for_bulk_import.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")
    print(f"  total unique products: {len(by_key):,}")
    print(f"  classified (cat set):  {classified:,}  ({classified*100//len(by_key)}%)")
    print(f"  needs review:          {review:,}  ({review*100//len(by_key)}%)")


if __name__ == "__main__":
    main()
